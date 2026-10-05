"""Turn raw AI text into HTML preview, .docx and .pdf. (docx/fpdf are imported lazily.)"""
import html
import io
import re

from config import FONT_BOLD, FONT_REGULAR, FOOTER_TEXT, LOGO_PATH

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u00a0": " ", "\u200b": "",
}
_HEADING_RE = re.compile(r"^\d+\.\s+[^.]{2,80}:?$")      # "1. Services:" (no inner full stop)
TERMS_TITLE = "Summary of Key Terms"


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def sanitize_text(text: str) -> str:
    """Remove typographic quotes and leftover markdown so every format looks clean."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)   # markdown headings
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"^\s*\*\s+", "- ", text, flags=re.MULTILINE)          # "* item" -> "- item"
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_terms(terms: str) -> list[str]:
    """Split the 'Terms & Conditions' box on semicolons / new lines."""
    return [t.strip(" \t-•") for t in re.split(r"[;\n]", terms or "") if t.strip(" \t-•")]


def is_heading(line: str) -> bool:
    return bool(_HEADING_RE.match(line.strip()))


def split_title(text: str, doc_type: str) -> tuple[str, list[str]]:
    """Return (title, remaining lines). Uses the AI's first line if it looks like a title."""
    lines = [ln.rstrip() for ln in text.split("\n")]
    first = next((ln for ln in lines if ln.strip()), "")
    if (first and len(first) <= 80 and not first.endswith((".", ":", ";"))
            and not re.match(r"^\d+\.", first) and len(lines) > 1):
        idx = lines.index(first)
        return first.strip(), lines[idx + 1:]
    return doc_type.strip().title() or "Legal Document", lines


def safe_filename(name: str, ext: str) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", (name or "document").lower()).strip("_") or "document"
    return f"{base}.{ext}"


# --------------------------------------------------------------------------
# HTML preview (escaped - AI text is never trusted as HTML)
# --------------------------------------------------------------------------
def format_html_preview(text: str, doc_type: str = "") -> str:
    title, lines = split_title(sanitize_text(text), doc_type)
    parts = [f"<h3 style='text-align:center;margin:0 0 12px'>{html.escape(title)}</h3>"]
    for line in lines:
        line = line.strip()
        if not line:
            continue
        safe = html.escape(line)
        if is_heading(line):
            parts.append(f"<p style='margin:14px 0 4px'><b>{safe}</b></p>")
        elif line.startswith("- "):
            parts.append(f"<p style='margin:2px 0 2px 18px'>&bull; {html.escape(line[2:])}</p>")
        else:
            parts.append(f"<p style='margin:4px 0'>{safe}</p>")
    return "\n".join(parts)


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    title, lines = split_title(sanitize_text(text), doc_type)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(2.2))

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    term_list = parse_terms(terms)
    if term_list:
        h = doc.add_paragraph()
        h.add_run(TERMS_TITLE).bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "#"
        table.rows[0].cells[1].text = "Term"
        for c in table.rows[0].cells:
            for r in c.paragraphs[0].runs:
                r.bold = True
        for i, term in enumerate(term_list, start=1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = term
        doc.add_paragraph()

    for line in lines:
        line = line.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        if is_heading(line):
            p.add_run(line).bold = True
            p.paragraph_format.space_before = Pt(10)
        elif line.startswith("- "):
            p.style = "List Bullet"
            p.add_run(line[2:])
        else:
            p.add_run(line)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(FOOTER_TEXT)
    fr.italic = True
    fr.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos
    from PIL import Image

    title, lines = split_title(sanitize_text(text), doc_type)
    logo_w = 45
    logo_h = 0
    if LOGO_PATH.exists():
        with Image.open(LOGO_PATH) as im:
            logo_h = logo_w * im.height / im.width

    class LegalPDF(FPDF):
        def header(self):                       # logo + title on every page
            top = 8
            if logo_h:
                self.image(str(LOGO_PATH), x=(self.w - logo_w) / 2, y=top, w=logo_w)
            self.set_y(top + logo_h + 3)
            self.set_font("DejaVu", "B", 12)
            self.multi_cell(0, 7, title, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.ln(4)

        def footer(self):
            self.set_y(-15)
            self.set_font("DejaVu", "", 8)
            self.cell(0, 8, FOOTER_TEXT, align="C")

    pdf = LegalPDF()
    pdf.add_font("DejaVu", "", str(FONT_REGULAR))
    pdf.add_font("DejaVu", "B", str(FONT_BOLD))
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    def write(style: str, size: int, content: str, gap: float = 2):
        pdf.set_font("DejaVu", style, size)
        # new_x/new_y are required in fpdf2, otherwise the next multi_cell(0, ...) fails
        pdf.multi_cell(0, 6, content, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(gap)

    term_list = parse_terms(terms)
    if term_list:
        write("B", 11, TERMS_TITLE, gap=1)
        for term in term_list:
            write("", 10, f"-  {term}", gap=0.5)
        pdf.ln(4)

    for line in lines:
        line = line.strip()
        if not line:
            continue
        if is_heading(line):
            pdf.ln(2)
            write("B", 11, line, gap=1)
        elif line.startswith("- "):
            write("", 10, f"-  {line[2:]}", gap=0.5)
        else:
            write("", 10, line)

    return bytes(pdf.output())
